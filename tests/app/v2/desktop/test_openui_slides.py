
from typing import TYPE_CHECKING, cast

if TYPE_CHECKING:
    from app.v2.desktop.backend.service import DesktopV2Service

from pptx.shapes.autoshape import Shape

import copy
import json
from io import BytesIO
from pathlib import Path
from unittest.mock import AsyncMock

import pytest
from fastapi.testclient import TestClient
from pptx import Presentation

from app.v2.desktop.backend.openui import OpenUiParser, validate_tree
from app.v2.desktop.backend.openui_slides import export_deck, slide_slots
from app.v2.desktop.backend import app as app_module
from tests.app.v2.desktop.test_app_lifecycle import _Service, AUTH_HEADERS, AUTH_TOKEN


@pytest.fixture
def deck():
    fixtures = json.loads(
        Path("app/v2/desktop/test/fixtures/openui_gallery.json").read_text()
    )
    return next(f["surface"]["tree"] for f in fixtures if f["name"] == "slides")


@pytest.mark.asyncio
async def test_slide_text_streams_and_export_contains_editable_objects(deck):
    parser = OpenUiParser()
    try:
        code = 'root = Deck("Show", [Slide("Title", [Text("Hel'
        partial = await parser.parse(code, complete=False)
        assert (
            partial["props"]["children"][0]["props"]["children"][0]["props"]["content"]
            == "Hel"
        )
        assert (
            await parser.parse(code + 'lo")], "content", "blue")])', complete=True)
        )["component"] == "Deck"
    finally:
        await parser.close()
    for dark in (False, True):
        prs = Presentation(BytesIO(export_deck(deck, dark=dark)))
        assert len(prs.slides) == 3
        assert cast('int', prs.slide_width) / cast('int', prs.slide_height) == pytest.approx(16 / 9)
        assert any(s.has_chart for slide in prs.slides for s in slide.shapes)
        assert any(s.has_table for slide in prs.slides for s in slide.shapes)
        assert any(
            "让想法成为作品" in cast(Shape, s).text
            for slide in prs.slides
            for s in slide.shapes
            if s.has_text_frame
        )
        assert not any(
            s.shape_type == 13 for slide in prs.slides for s in slide.shapes
        )  # no screenshot pictures
        chart = next(
            getattr(s, 'chart') for slide in prs.slides for s in slide.shapes if s.has_chart
        )
        assert list(chart.series[0].values) == [72, 96, 128]
        assert chart.part.chart_workbook.xlsx_part.blob.startswith(b"PK")


def test_slide_layout_limits_prevent_forms_and_oversized_data(deck):
    assert len(slide_slots("split", 4)) == 4
    for box in slide_slots("split", 4):
        x, y, w, h = box
        assert x >= 64 and y >= 206 and x + w <= 1216 and y + h <= 646
    bad = copy.deepcopy(deck)
    bad["props"]["children"][0]["props"]["children"][0]["props"]["content"] = "x" * 241
    with pytest.raises(ValueError):
        validate_tree(bad)
    bad = copy.deepcopy(deck)
    bad["props"]["children"][1]["props"]["children"][0] = {
        "component": "Button",
        "props": {"action": "x", "label": "x"},
    }
    with pytest.raises(ValueError, match="Unsupported slide"):
        validate_tree(bad)
    bad = copy.deepcopy(deck)
    bad["props"]["children"][1]["props"]["children"][1]["props"]["series"][0]["values"][
        0
    ] = "field"
    with pytest.raises(ValueError, match="literal"):
        validate_tree(bad)


def test_export_endpoint_auth_session_revision_and_payload(deck):
    service = _Service()
    setattr(service, 'openui_surfaces', AsyncMock(
        return_value=[{"surface_id": "deck1", "revision": 2, "tree": deck}]
    ))
    app = app_module.create_app(service=cast('DesktopV2Service | None', service), auth_token=AUTH_TOKEN)
    with TestClient(app) as client:
        path = "/api/v2/sessions/s1/openui/deck1/pptx"
        assert client.get(path + "?revision=2").status_code == 401
        assert client.get(path + "?revision=1", headers=AUTH_HEADERS).status_code == 409
        assert (
            client.get(
                path.replace("deck1", "missing") + "?revision=2", headers=AUTH_HEADERS
            ).status_code
            == 404
        )
        response = client.get(path + "?revision=2&dark=true", headers=AUTH_HEADERS)
        assert response.status_code == 200
        import base64

        assert (
            len(
                Presentation(
                    BytesIO(base64.b64decode(response.json()["data"]["base64"]))
                ).slides
            )
            == 3
        )
    getattr(service, 'openui_surfaces').assert_awaited_with("s1", "default_user")


@pytest.mark.parametrize("kind", ["line", "bar", "pie"])
def test_all_presentation_chart_types_export_editable(deck, kind):
    tree = copy.deepcopy(deck)
    tree["props"]["children"][1]["props"]["children"][1]["props"]["type"] = kind
    prs = Presentation(BytesIO(export_deck(tree)))
    chart = next(getattr(s, 'chart') for slide in prs.slides for s in slide.shapes if s.has_chart)
    assert list(chart.series[0].values) == [72, 96, 128]
